# Reads test-case questions from the benchmark xlsx and writes results out to
# a new results workbook (the input file is never modified).

from __future__ import annotations

from pathlib import Path

import openpyxl


def load_test_cases(xlsx_path: str | Path, *, sheet_name: str, question_column: str) -> list[str]:
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb[sheet_name] if sheet_name in wb.sheetnames else wb.worksheets[0]

    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    header = [str(c).strip() if c is not None else "" for c in rows[0]]
    try:
        col_idx = header.index(question_column)
    except ValueError:
        col_idx = 0  # fall back to first column

    questions = []
    for row in rows[1:]:
        if col_idx >= len(row):
            continue
        value = row[col_idx]
        if value is None:
            continue
        text = str(value).strip()
        if text:
            questions.append(text)
    return questions


def write_summary_xlsx(results: list[dict], out_path: str | Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "results"

    columns = [
        "test_index",
        "condition",
        "question",
        "status",
        "duration_seconds",
        "total_tokens",
        "input_tokens",
        "output_tokens",
        "main_total_tokens",
        "a2ui_repair_total_tokens",
        "a2ui_repair_call_count",
        "a2ui_retry_plain_total_tokens",
        "a2ui_retry_plain_call_count",
        "error",
    ]
    ws.append(columns)
    for row in results:
        usage_by_purpose = row.get("usage_by_purpose", {})
        main = usage_by_purpose.get("main", {})
        repair = usage_by_purpose.get("a2ui_repair", {})
        retry = usage_by_purpose.get("a2ui_retry_plain", {})
        ws.append([
            row.get("test_index"),
            row.get("condition"),
            row.get("question"),
            row.get("status"),
            round(row.get("duration_seconds", 0), 1),
            row.get("total_tokens", 0),
            row.get("input_tokens", 0),
            row.get("output_tokens", 0),
            main.get("total_tokens", 0),
            repair.get("total_tokens", 0),
            repair.get("call_count", 0),
            retry.get("total_tokens", 0),
            retry.get("call_count", 0),
            row.get("error", ""),
        ])

    # Comparison sheet: with_a2ui vs without_a2ui totals per test_index.
    by_index: dict[int, dict[str, dict]] = {}
    for row in results:
        by_index.setdefault(row["test_index"], {})[row["condition"]] = row

    cmp_ws = wb.create_sheet("comparison")
    cmp_ws.append([
        "test_index", "question",
        "with_a2ui_total_tokens", "without_a2ui_total_tokens", "delta_tokens", "delta_pct",
        "with_a2ui_status", "without_a2ui_status", "clean_pair",
    ])
    clean_pairs = []
    for idx in sorted(by_index):
        entry = by_index[idx]
        with_row = entry.get("with_a2ui", {})
        without_row = entry.get("without_a2ui", {})
        with_tokens = with_row.get("total_tokens", 0)
        without_tokens = without_row.get("total_tokens", 0)
        delta = with_tokens - without_tokens
        delta_pct = (delta / without_tokens * 100) if without_tokens else None
        # "Clean" = both conditions completed without hitting the per-test
        # timeout or erroring - only these are directly comparable, since a
        # timeout/error run's token count reflects however much happened to
        # be captured before it was cut off, not the true cost of that turn.
        is_clean = with_row.get("status") == "ok" and without_row.get("status") == "ok"
        if is_clean:
            clean_pairs.append((with_tokens, without_tokens))
        cmp_ws.append([
            idx,
            with_row.get("question") or without_row.get("question"),
            with_tokens,
            without_tokens,
            delta,
            round(delta_pct, 1) if delta_pct is not None else "",
            with_row.get("status", ""),
            without_row.get("status", ""),
            is_clean,
        ])

    summary_ws = wb.create_sheet("summary")
    summary_ws.append(["metric", "value"])
    summary_ws.append(["total_test_cases", len(by_index)])
    summary_ws.append(["clean_pairs (both conditions status=ok)", len(clean_pairs)])
    summary_ws.append(["excluded_pairs (timeout/error in either condition)", len(by_index) - len(clean_pairs)])
    if clean_pairs:
        with_clean_total = sum(p[0] for p in clean_pairs)
        without_clean_total = sum(p[1] for p in clean_pairs)
        summary_ws.append(["clean_with_a2ui_total_tokens", with_clean_total])
        summary_ws.append(["clean_without_a2ui_total_tokens", without_clean_total])
        summary_ws.append(["clean_overhead_tokens", with_clean_total - without_clean_total])
        summary_ws.append([
            "clean_overhead_pct",
            round((with_clean_total - without_clean_total) / without_clean_total * 100, 1)
            if without_clean_total else "",
        ])
        summary_ws.append(["clean_avg_with_a2ui_tokens_per_test", round(with_clean_total / len(clean_pairs))])
        summary_ws.append(["clean_avg_without_a2ui_tokens_per_test", round(without_clean_total / len(clean_pairs))])
    summary_ws.append([
        "note",
        "Only clean_pairs (status=ok in both conditions) are directly comparable. "
        "Timeout/error runs are excluded from clean_* metrics because their token "
        "count reflects only what was captured before the run was cut off, not the "
        "true cost of that turn - see comparison sheet's clean_pair column and "
        "README's Timeouts and failure handling section.",
    ])

    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
